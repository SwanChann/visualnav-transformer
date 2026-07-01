# -*- coding: utf-8 -*-
from __future__ import annotations

from pathlib import Path

from pptx import Presentation


PPT_PATH = Path(__file__).resolve().parent / "陈样答辩.pptx"

NOTES = [
    "各位老师好，我是陈样。我的题目是《基于动作扩散策略的四足机器人视觉目标导航》。课题以 NoMaD 为高层视觉策略，以 Lite3 四足机器人为平台，目标是让机器人根据当前图像和目标图像，在真实环境中闭环到达目标。下面按背景、算法、系统和实验结果汇报。",
    "这一页说明视觉目标导航的任务形态。机器人从当前视觉观察和目标图像出发，经策略推理生成下一步动作，再由四足平台执行。本文关注的是把“看见目标”变成“走到目标”：用 NoMaD 生成局部动作，并在 Lite3 上形成稳定闭环。",
    "选择扩散策略，是因为它能建模多模态动作分布，适合岔路选择、绕行和目标不确定的场景；选择四足机器人，是因为它更适合复杂真实环境。现有工作较少把扩散视觉导航落到四足真机，因此本文聚焦实时推理、动作接口和运动稳定控制。",
    "本文的核心问题是：在算力受限、动力学复杂、环境变化明显的条件下，让四足机器人依据目标图像稳定导航。对应三类挑战：算法上扩散采样慢；系统上要把局部轨迹转成 Lite3 速度指令；实验上要从仿真走到真实环境。",
    "总体路线是四步：先梳理视觉目标导航、扩散模型和 NoMaD；再构建统一推理模块，比较 DDIM、CFG、TTS 和编码器；随后搭建分层闭环系统并做 MuJoCo 验证；最后迁移到 Orin-Lite3 真机。主线是算法优化、系统闭环、真机验证逐步推进。",
    "算法基线是 NoMaD。它输入历史观测、目标图像和目标掩码，视觉编码器提取条件特征；距离预测器估计拓扑节点远近；扩散策略头生成局部航点。本文保留这个框架，重点优化采样速度、候选筛选和四足适配。模型约 19.05M 参数，输出 8 条轨迹、每条 8 个航点。",
    "这一页展示 NoMaD 的两种模式。导航模式下，模型根据目标图像生成朝向目标的候选轨迹；探索模式下，目标掩码置空，策略依据历史观察生成前进动作。后续真机实验也分别验证了目标导航和无目标探索。",
    "本文把四类优化放进统一推理模块：DDIM 减少扩散去噪步数，主要解决速度；CFG 调节目标条件引导，但会增加开销；TTS 生成多候选并筛选；编码器替换比较不同视觉表征。最终默认配置是 EfficientNet-B0、DDIM-2、CFG=0、TTS-8。",
    "离线实验我简要带过，重点看结论。DDIM 加速实验显示，DDIM-2 相对 DDPM-10 约有 4.84 倍纯采样加速，同时轨迹质量仍可接受；DDIM-1 虽更快，但偏离基线明显，因此不采用。",
    "轨迹可视化用于直观看不同 DDIM 步数的动作形态。除 DDIM-1 外，多数配置轨迹趋势与基线接近。因此本文选择 DDIM-2，不追求最少步数，而是在速度和稳定性之间取折中。",
    "CFG 的作用是增强目标图像的条件引导，但需要有条件和无条件两次分支推理，开销接近翻倍。离线收益不稳定，结合真机实时性，本文把 CFG 作为可调参数，而不是默认开启。",
    "TTS 的思想是一次生成多条候选轨迹，再用评分选择更合适的动作。实验中 TTS-8 在质量和时延之间较平衡，TTS-16 提升有限，TTS-32 边际收益下降，所以真机默认使用 TTS-8。",
    "编码器实验说明，不同视觉表征各有优势：ResNet-50 的多样性和前向进展较高，EfficientNet-B0 更平滑。但真机部署不能只看离线指标，还要考虑真实相机分布、Orin 时延和安全约束，因此默认仍采用 EfficientNet-B0。",
    "离线部分总结为三点：速度瓶颈主要来自扩散采样，DDIM-2 是核心加速；CFG 开销较大，默认关闭；TTS-8 能低成本保留候选多样性。后续系统和真机实验都使用 EfficientNet-B0、DDIM-2、CFG=0、TTS-8。",
    "算法要进入机器人，还需要系统闭环。本文分为四层：推理层输出轨迹和距离；任务组织层维护目标节点与状态机；中层控制层把轨迹映射为线速度和角速度；平台执行层对接 MuJoCo 或 Lite3 UDP。这样仿真和真机可以复用同一任务逻辑。",
    "MuJoCo 的作用是健康检查，不替代真机结论。这里验证状态机、PD 映射和平台接口能否形成完整闭环。DDIM、CFG 和不同编码器配置都能接入仿真；通过后，高层策略保持不变，只把平台执行层切换到 Lite3 UDP 链路。",
    "真机系统由两条链路组成。水平推理链从相机图像进入 NoMaD，经距离预测、扩散策略和 PD 映射得到速度指令；垂直控制链由 Orin 主机、四足接口和底层运控组成。为保证安全，速度刷新 25Hz，心跳 5Hz，最大线速度 0.12m/s，最大角速度 0.35rad/s。",
    "真实导航使用预采集拓扑地图。运行前沿路线采集有序视觉节点；运行时只在当前节点附近建立正负 4 的滑动窗口，用距离预测器选择窗口内目标。这样把长距离任务拆成连续短目标，避免一次跳到远端节点。",
    "真机结果表明，四组配置在室内长廊、有障碍连廊和室外斜坡中均 100% 成功，在草坪场景失败。最快方案 DDIM-2、CFG=0、TTS-8 把闭环频率从 DDPM 基线的 2.89Hz 提升到 6.80Hz，约为 2.4 倍，说明离线加速转化成了真机收益。",
    "这一页展示探索和室外泛化。无目标探索时，目标掩码置空，策略根据历史观察前进；室外零样本测试直接迁移到斜坡、草地和强光环境。简单斜坡可以运行，但草地、开阔区域和强反光会使视觉表征不稳定，说明系统仍受训练分布限制。",
    "最后总结。本文完成三件事：构建统一 NoMaD 推理模块，比较 DDIM、CFG、TTS 和编码器；搭建分层闭环系统，实现仿真与 Lite3 真机接口复用；完成 MuJoCo、室内导航、无目标探索和室外泛化验证。后续将补充四足视角数据，加强控制融合，并评估 TensorRT 和量化部署。",
    "以上就是我的毕业设计汇报。本文把动作扩散视觉导航方法推进到 Lite3 四足机器人闭环系统，并完成仿真和真机验证。感谢各位老师聆听，请批评指正。",
]


def main() -> None:
    prs = Presentation(str(PPT_PATH))
    if len(prs.slides) != len(NOTES):
        raise RuntimeError(f"slide count {len(prs.slides)} != notes count {len(NOTES)}")

    for slide, note in zip(prs.slides, NOTES):
        frame = slide.notes_slide.notes_text_frame
        frame.clear()
        frame.text = note

    prs.save(str(PPT_PATH))

    reread = Presentation(str(PPT_PATH))
    notes = [
        slide.notes_slide.notes_text_frame.text.strip()
        if slide.has_notes_slide
        else ""
        for slide in reread.slides
    ]
    bad = [(idx + 1, text.count("?"), text[:20]) for idx, text in enumerate(notes) if text.count("?") > 2]
    chinese_count = sum(sum("\u4e00" <= ch <= "\u9fff" for ch in text) for text in notes)
    print(f"saved={PPT_PATH}")
    print(f"slides={len(notes)} notes={sum(bool(text) for text in notes)} chinese_chars={chinese_count}")
    print(f"bad_qmark_slides={bad}")


if __name__ == "__main__":
    main()
